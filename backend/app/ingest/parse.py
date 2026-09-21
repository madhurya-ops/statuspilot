"""Turn an uploaded file into plain transcript text.

Supported: .txt, .docx, .vtt, .srt. Everything here is pure text wrangling — no
network, no LLM — so it is fully covered by fast unit tests.
"""

import io
import re

from fastapi import HTTPException, status

ALLOWED_EXTENSIONS = {".txt", ".docx", ".vtt", ".srt"}

# 00:01:02.345 --> 00:01:07.890  (VTT) and 00:01:02,345 --> 00:01:07,890 (SRT),
# optionally followed by VTT cue settings such as "align:start position:10%".
_TIMESTAMP = re.compile(
    r"^\d{1,2}:\d{2}(:\d{2})?[.,]\d{1,3}\s*-->\s*\d{1,2}:\d{2}(:\d{2})?[.,]\d{1,3}"
)
_CUE_NUMBER = re.compile(r"^\d+$")
# VTT voice span: <v Priya Nair>text</v>
_VOICE = re.compile(r"<v\.?[^\s>]*\s+([^>]+)>(.*?)(?:</v>)?$", re.IGNORECASE | re.DOTALL)
_TAG = re.compile(r"</?[^>]+>")


def decode_text(data: bytes) -> str:
    """Decode bytes as UTF-8, falling back to latin-1, which cannot fail."""
    try:
        return data.decode("utf-8")
    except UnicodeDecodeError:
        return data.decode("latin-1")


def parse_docx(data: bytes) -> str:
    """Extract paragraph text from a .docx, including text inside tables."""
    import docx  # imported lazily: only this path needs python-docx

    try:
        document = docx.Document(io.BytesIO(data))
    except Exception as err:
        raise HTTPException(
            status_code=status.HTTP_400_BAD_REQUEST,
            detail="That .docx could not be read. It may be corrupt or password-protected.",
        ) from err

    parts = [p.text for p in document.paragraphs]
    for table in document.tables:
        for row in table.rows:
            cells = [c.text.strip() for c in row.cells if c.text.strip()]
            if cells:
                parts.append(" | ".join(cells))
    return "\n".join(parts)


def parse_captions(text: str) -> str:
    """Flatten a .vtt or .srt caption file into `Speaker: text` lines.

    Caption formats split a single sentence across cues, so consecutive cues from
    the same speaker are merged. Without that, a spoken commitment lands as three
    fragments and neither the LLM nor Jev sees a whole item.
    """
    blocks: list[tuple[str | None, str]] = []

    for raw in text.replace("\r\n", "\n").replace("\r", "\n").split("\n"):
        line = raw.strip()
        if not line:
            continue
        upper = line.upper()
        if upper.startswith("WEBVTT") or upper.startswith("NOTE") or upper.startswith("STYLE"):
            continue
        if _TIMESTAMP.match(line) or _CUE_NUMBER.match(line):
            continue

        speaker: str | None = None
        voice = _VOICE.match(line)
        if voice:
            speaker = voice.group(1).strip()
            line = voice.group(2).strip()
        line = _TAG.sub("", line).strip()
        if not line:
            continue

        if speaker is None:
            # Fall back to a plain "Name: text" prefix inside the cue.
            from app.ingest.lines import split_speaker

            speaker, line = split_speaker(line)

        if blocks and blocks[-1][0] == speaker:
            blocks[-1] = (speaker, f"{blocks[-1][1]} {line}".strip())
        else:
            blocks.append((speaker, line))

    return "\n".join(f"{s}: {t}" if s else t for s, t in blocks)


def parse_upload(filename: str, data: bytes) -> str:
    """Dispatch on file extension. Raises 400 for anything not allow-listed."""
    suffix = "." + filename.rsplit(".", 1)[-1].lower() if "." in filename else ""
    if suffix not in ALLOWED_EXTENSIONS:
        raise HTTPException(
            status_code=status.HTTP_400_BAD_REQUEST,
            detail=f"Unsupported file type. Allowed: {', '.join(sorted(ALLOWED_EXTENSIONS))}.",
        )
    if suffix == ".docx":
        return parse_docx(data)
    text = decode_text(data)
    if suffix in {".vtt", ".srt"}:
        return parse_captions(text)
    return text
