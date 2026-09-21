"""Status report + minutes as a .docx.

Built from the structured items, not by parsing the generated markdown, so the
action-item table cannot contain anything the data does not.
"""

import io
import re

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt, RGBColor

from app.models import Documents

RAG_COLOUR = {
    "Red": RGBColor(0xB4, 0x23, 0x18),
    "Amber": RGBColor(0xB5, 0x47, 0x08),
    "Green": RGBColor(0x06, 0x76, 0x47),
}


def _markdown_into(document: Document, source: str) -> None:
    """Render the narrative subset: headings, bullets, paragraphs, bold runs."""
    for raw in source.split("\n"):
        line = raw.rstrip()
        if not line.strip():
            continue
        heading = re.match(r"^(#{1,4})\s+(.*)$", line)
        if heading:
            document.add_heading(heading.group(2), level=min(3, len(heading.group(1)) + 1))
            continue
        if re.match(r"^\s*[-*+]\s+", line):
            text = re.sub(r"^\s*[-*+]\s+", "", line)
            _runs(document.add_paragraph(style="List Bullet"), text)
            continue
        if line.strip().startswith("|"):
            continue  # tables are added from the data, never from prose
        _runs(document.add_paragraph(), line)


def _runs(paragraph, text: str) -> None:
    for part in re.split(r"(\*\*[^*]+\*\*)", text):
        if part.startswith("**") and part.endswith("**"):
            paragraph.add_run(part[2:-2]).bold = True
        elif part:
            paragraph.add_run(part)


def build_docx(documents: Documents, project: str, rag: str) -> bytes:
    document = Document()
    style = document.styles["Normal"]
    style.font.name = "Calibri"
    style.font.size = Pt(11)

    title = document.add_heading(project or "Project status report", level=0)
    title.alignment = WD_ALIGN_PARAGRAPH.LEFT

    status = document.add_paragraph()
    status.add_run("Overall status: ").bold = True
    run = status.add_run(rag)
    run.bold = True
    run.font.color.rgb = RAG_COLOUR.get(rag, RGBColor(0x23, 0x24, 0x33))

    _markdown_into(document, documents.status_report_markdown)

    if documents.action_items:
        document.add_heading("Action items", level=1)
        table = document.add_table(rows=1, cols=4)
        table.style = "Light Grid Accent 1"
        for cell, label in zip(
            table.rows[0].cells, ("Action", "Owner", "Due", "Severity"), strict=True
        ):
            cell.paragraphs[0].add_run(label).bold = True
        for item in documents.action_items:
            cells = table.add_row().cells
            cells[0].text = item.text
            cells[1].text = item.owner or "Not assigned"
            cells[2].text = item.due_date or "No date"
            cells[3].text = item.severity

    for section, items in documents.raid_log.items():
        if not items:
            continue
        document.add_heading(section, level=1)
        for item in items:
            paragraph = document.add_paragraph(style="List Bullet")
            paragraph.add_run(f"[{item.severity}] ").bold = True
            paragraph.add_run(item.text)

    document.add_page_break()
    document.add_heading("Minutes of meeting", level=1)
    _markdown_into(document, documents.mom_markdown)

    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()
