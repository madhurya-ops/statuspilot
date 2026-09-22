"""RAID log + action items as a workbook — the format PMs actually live in."""

import io

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from app.export.audience import prepare
from app.models import Documents

HEADER_FILL = PatternFill("solid", fgColor="4F46E5")
HEADER_FONT = Font(color="FFFFFF", bold=True, size=11)
SEVERITY_FILL = {
    "High": PatternFill("solid", fgColor="FEF3F2"),
    "Medium": PatternFill("solid", fgColor="FFFAEB"),
    "Low": PatternFill("solid", fgColor="F5F5F8"),
}

COLUMNS = [
    "Item", "Owner", "Owner source", "Due", "Due source",
    "Severity", "Score", "Audience", "Lines",
]
WIDTHS = [62, 18, 14, 16, 12, 10, 8, 14, 12]


def _sheet(workbook: Workbook, title: str, items) -> None:
    sheet = workbook.create_sheet(title[:31])
    sheet.append(COLUMNS)
    for index, width in enumerate(WIDTHS, start=1):
        sheet.column_dimensions[get_column_letter(index)].width = width
    for cell in sheet[1]:
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = Alignment(vertical="center")
    sheet.freeze_panes = "A2"

    for item in prepare(items, "internal"):
        sheet.append([
            item.text,
            item.owner or "Not assigned",
            item.owner_status.replace("_", " "),
            item.due_date or "No date",
            item.due_status.replace("_", " "),
            item.severity,
            # The raw float, kept so a PM can sort by real priority rather than by
            # a three-way band that ties 1.52 with 2.0.
            round(item.severity_value, 2),
            "Client" if item.audience == "client_safe" else "Internal",
            ", ".join(f"L{n}" for n in item.source_lines),
        ])
        row = sheet[sheet.max_row]
        fill = SEVERITY_FILL.get(item.severity)
        if fill:
            row[5].fill = fill
        row[0].alignment = Alignment(wrap_text=True, vertical="top")

    if not items:
        sheet.append(["Nothing recorded in this category."])


def build_xlsx(documents: Documents) -> bytes:
    """The internal working file: every item, with an Audience column saying which
    are cleared for the client. Provenance tags are reconciled with their values so a
    row never reads `Owner: Raj Menon / Owner source: not specified`."""
    workbook = Workbook()
    workbook.remove(workbook.active)
    _sheet(workbook, "Action Items", documents.action_items)
    for section in ("Risks", "Assumptions", "Issues", "Dependencies"):
        _sheet(workbook, section, documents.raid_log.get(section, []))
    buffer = io.BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()
