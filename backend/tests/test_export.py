"""Phase 8: every generated file must re-open in the library that made it."""

import io

import pytest

from app.models import ApprovedItem, Documents
from app.routers.export import safe_filename


def _item(**kw):
    base = dict(
        id="c1",
        text="Fix the orphaned document reference — before UAT",
        kind="action_item",
        severity="High",
        severity_value=1.92,
        audience="client_safe",
        owner="Marcus Webb",
        owner_status="stated",
        due_date="Thursday",
        due_status="stated",
        source_lines=[19, 20],
    )
    base.update(kw)
    return ApprovedItem(**base)


@pytest.fixture
def documents():
    return Documents(
        mom_markdown="## Attendees\nPriya Nair\n\n## Decisions\n- Ship it",
        status_report_markdown=(
            "## Overall status: Red\n\nThe date **moved** — badly.\n\n"
            "## Next steps\n- Complete the fix by 27 March"
        ),
        action_items=[
            _item(),
            _item(id="c2", severity="Low", severity_value=0.2, owner=None, due_date=None),
        ],
        raid_log={
            "Risks": [_item(id="c3", kind="risk", audience="internal_only")],
            "Assumptions": [],
            "Issues": [],
            "Dependencies": [],
        },
    )


class TestDocx:
    def test_reopens_in_python_docx(self, documents):
        from docx import Document

        from app.export.docx_export import build_docx

        blob = build_docx(documents, "Contoso Insurance", "Red")
        reopened = Document(io.BytesIO(blob))
        text = "\n".join(p.text for p in reopened.paragraphs)
        assert "Contoso Insurance" in text
        assert "Red" in text
        assert reopened.tables, "the action-item table must be present"

    def test_the_action_table_comes_from_the_data(self, documents):
        from docx import Document

        from app.export.docx_export import build_docx

        reopened = Document(io.BytesIO(build_docx(documents, "P", "Red")))
        rows = reopened.tables[0].rows
        assert len(rows) == 1 + len(documents.action_items)
        assert rows[1].cells[1].text == "Marcus Webb"
        assert rows[2].cells[1].text == "Not assigned"

    def test_markdown_tables_are_never_carried_into_the_document(self, documents):
        from docx import Document

        from app.export.docx_export import build_docx

        documents.status_report_markdown += "\n\n| a | b |\n|---|---|\n| 1 | 2 |\n"
        reopened = Document(io.BytesIO(build_docx(documents, "P", "Red")))
        text = "\n".join(p.text for p in reopened.paragraphs)
        assert "| a | b |" not in text


class TestXlsx:
    def test_reopens_with_all_five_sheets(self, documents):
        from openpyxl import load_workbook

        from app.export.xlsx_export import build_xlsx

        book = load_workbook(io.BytesIO(build_xlsx(documents)))
        assert book.sheetnames == ["Action Items", "Risks", "Assumptions", "Issues", "Dependencies"]

    def test_header_is_frozen_and_the_raw_score_is_present(self, documents):
        from openpyxl import load_workbook

        from app.export.xlsx_export import build_xlsx

        sheet = load_workbook(io.BytesIO(build_xlsx(documents)))["Action Items"]
        assert sheet.freeze_panes == "A2"
        header = [c.value for c in sheet[1]]
        assert "Score" in header
        # The float is exported so a PM can sort by real priority, not by a band
        # that ties 1.52 with 2.0.
        assert sheet.cell(row=2, column=header.index("Score") + 1).value == 1.92

    def test_an_empty_category_says_so(self, documents):
        from openpyxl import load_workbook

        from app.export.xlsx_export import build_xlsx

        sheet = load_workbook(io.BytesIO(build_xlsx(documents)))["Issues"]
        assert sheet.cell(row=2, column=1).value


class TestPdf:
    def test_produces_a_valid_pdf(self, documents):
        from app.export.pdf_export import build_pdf

        blob = build_pdf(documents, "Contoso Insurance", "Red")
        assert blob.startswith(b"%PDF-")
        assert blob.rstrip().endswith(b"%%EOF")

    def test_survives_long_text_and_many_items(self, documents):
        """A drifting cursor once made fpdf2 raise "Not enough horizontal space"."""
        from app.export.pdf_export import build_pdf

        documents.action_items = [_item(text="A very long action item. " * 12)] * 12
        assert build_pdf(documents, "A very long project name " * 3, "Amber").startswith(b"%PDF-")

    def test_non_latin1_characters_do_not_crash_it(self, documents):
        from app.export.pdf_export import build_pdf

        documents.status_report_markdown = "## Status\n\nZoë — café, naïve… “quoted”"
        assert build_pdf(documents, "Zoë Café", "Green").startswith(b"%PDF-")


class TestFilename:
    @pytest.mark.parametrize(
        "project,expected_stem",
        [
            ("Contoso Insurance", "Contoso-Insurance"),
            ("../../etc/passwd", "etc-passwd"),
            ('a"; rm -rf /', "a-rm-rf"),
            (None, "Project"),
            ("", "Project"),
        ],
    )
    def test_is_sanitised(self, project, expected_stem):
        """The project name is user input and lands in a response header."""
        name = safe_filename(project, "pdf")
        assert name.startswith(f"StatusReport_{expected_stem}_")
        assert name.endswith(".pdf")
        assert "/" not in name and '"' not in name


class TestExportEndpoint:
    def _payload(self, documents):
        return {"documents": documents.model_dump(), "project_name": "Contoso", "rag": "Red"}

    def test_requires_the_access_code(self, client, documents):
        assert client.post("/api/export/pdf", json=self._payload(documents)).status_code == 401

    @pytest.mark.parametrize("fmt", ["docx", "xlsx", "pdf"])
    def test_each_format_downloads(self, client, auth, documents, fmt):
        r = client.post(f"/api/export/{fmt}", headers=auth, json=self._payload(documents))
        assert r.status_code == 200
        assert len(r.content) > 500
        assert "attachment" in r.headers["content-disposition"]
        assert fmt in r.headers["content-disposition"]

    def test_an_unknown_format_is_rejected(self, client, auth, documents):
        r = client.post("/api/export/exe", headers=auth, json=self._payload(documents))
        assert r.status_code == 400


class TestClientFacingExportsNeverLeak:
    """The PDF is the client-facing export and must not contain an internal item.

    This is not hypothetical: the PDF filtered its prose (which arrives pre-filtered
    from the generator) but appended `Documents.action_items` unfiltered, so internal
    remarks reached a file intended for a client.

    Asserted against the real committed cache rather than a fixture, because that is
    what the demo actually ships.
    """

    @pytest.fixture(params=["contoso-escalation", "northwind-sprint-review", "rough-standup-notes"])
    def cached(self, request):
        from app.pipeline.cache import load

        loaded = load(request.param)
        if loaded is None:
            pytest.skip(f"no cached run for {request.param}")
        return loaded["documents"]

    def _internal_texts(self, documents):
        every = list(documents.action_items)
        for items in documents.raid_log.values():
            every.extend(items)
        return [i.text for i in every if i.audience != "client_safe"]

    def test_no_internal_item_text_appears_in_the_pdf(self, cached):
        from app.export.pdf_export import _pdf_text, build_pdf

        blob = build_pdf(cached, "Contoso Insurance", "Red")
        # The PDF is written uncompressed precisely so this is checkable.
        body = blob.decode("latin-1")
        for text in self._internal_texts(cached):
            # Compare on a distinctive fragment, normalised the same way the PDF is.
            fragment = _pdf_text(text)[:48].strip()
            if len(fragment) < 20:
                continue
            assert fragment not in body, f"internal item leaked into the PDF: {fragment!r}"

    def test_the_pdf_still_contains_client_safe_material(self, cached):
        from app.export.pdf_export import build_pdf

        blob = build_pdf(cached, "Contoso Insurance", "Red")
        assert len(blob) > 800, "a filtered PDF should still carry the report"

    def test_the_docx_declares_itself_internal(self, cached):
        import io

        from docx import Document

        from app.export.docx_export import build_docx

        reopened = Document(io.BytesIO(build_docx(cached, "Contoso", "Red")))
        text = "\n".join(p.text for p in reopened.paragraphs)
        assert "INTERNAL" in text
        assert "not cleared for the client" in text


class TestProvenanceIsSelfConsistent:
    """A row must never read `Owner: Raj Menon / Owner source: not specified`.

    The value and the tag come from different models — the LLM extracts the owner,
    Jev judges whether the transcript explicitly names one — and they disagreed on 12
    rows of a single sample.
    """

    @pytest.mark.parametrize("sample", ["contoso-escalation", "northwind-sprint-review"])
    def test_no_contradictory_rows_reach_the_workbook(self, sample):
        import io

        from openpyxl import load_workbook

        from app.export.xlsx_export import build_xlsx
        from app.pipeline.cache import load

        loaded = load(sample)
        if loaded is None:
            pytest.skip(f"no cached run for {sample}")
        book = load_workbook(io.BytesIO(build_xlsx(loaded["documents"])))
        for sheet in book.worksheets:
            header = [c.value for c in sheet[1]]
            if "Owner" not in header:
                continue
            cols = {name: header.index(name) for name in header if name}
            for row in sheet.iter_rows(min_row=2, values_only=True):
                if not row or row[0] is None or len(row) < len(header):
                    continue
                owner, owner_src = row[cols["Owner"]], row[cols["Owner source"]]
                due, due_src = row[cols["Due"]], row[cols["Due source"]]
                has_owner = owner not in (None, "", "Not assigned")
                has_due = due not in (None, "", "No date")
                assert has_owner or owner_src == "not specified", (owner, owner_src)
                assert has_due or due_src == "not specified", (due, due_src)
                assert not has_owner or owner_src != "not specified", (owner, owner_src)
                assert not has_due or due_src != "not specified", (due, due_src)


class TestPdfCharacterHandling:
    def test_characters_outside_latin1_are_folded_not_dropped(self):
        from app.export.pdf_export import _pdf_text

        # U+2011 produced "go?live" in a real report, four times.
        assert _pdf_text("go‑live") == "go-live"
        assert _pdf_text("the vendor’s API") == "the vendor's API"
        assert _pdf_text("“quoted”") == '"quoted"'
        assert _pdf_text("a…b") == "a...b"
        assert _pdf_text("29 May – 5 June") == "29 May - 5 June"

    def test_accented_latin1_characters_survive_unchanged(self):
        from app.export.pdf_export import _pdf_text

        assert _pdf_text("Zoë café naïve") == "Zoë café naïve"

    def test_anything_genuinely_unrepresentable_degrades_rather_than_crashes(self):
        from app.export.pdf_export import _pdf_text

        out = _pdf_text("emoji 🙂 and CJK 日本語")
        assert "?" in out
        assert out.encode("latin-1")

    def test_a_report_full_of_awkward_characters_still_builds(self):
        from app.export.pdf_export import build_pdf

        documents = Documents(
            mom_markdown="## A",
            status_report_markdown="## Status\n\ngo‑live — the vendor’s “fix”… 🙂",
            action_items=[_item(text="Zoë to review the go‑live date – urgent")],
            raid_log={"Risks": [], "Assumptions": [], "Issues": [], "Dependencies": []},
        )
        assert build_pdf(documents, "Contoso — Migration", "Red").startswith(b"%PDF-")
